# -*- coding: utf-8 -*-
"""Reproduceerbare structuurcontrole; geen vervanging voor vakgroep-/proefcontrole."""
from pathlib import Path
import re
from zipfile import ZipFile
from docx import Document
from openpyxl import load_workbook
import pymupdf
from course_sheets import SHEETS, SHEET_MATERIAL
from course_guidance import G, FORMULAS, FORMULA_DISPLAY, SMART_EXAMPLES
from course_sheets import SHEET_SAFETY
from labs import LABOS
from answers import ANSWERS
from rubric_data import RUBRIC

ROOT=Path(__file__).resolve().parent.parent
assert len(SHEETS)==len(G)==len(SMART_EXAMPLES)==25 and len({title for _,_,title in SHEETS})==25
assert len(FORMULAS)>=20 and all(len(g.prior)==2 and all(g.transfer) for g in G)
assert not any(n in (6,7) for n,_,_ in SHEETS), 'Arbeid/momenten opnieuw opgenomen'
assert [(n,p) for n,p,_ in SHEETS[:6]]==[(1,(1,)),(1,(2,)),(2,(1,)),(2,(2,)),(3,(1,)),(3,(2,3))]
assert [p for n,p,_ in SHEETS if n==14]==[(1,),(2,),(3,)]
assert [p for n,p,_ in SHEETS if n==9]==[(1,),(3,),(2,)]
assert 'twintig' in ' '.join(LABOS[7]['delen'][0]['stappen']).lower()
assert 'geluidlab.html' in LABOS[12]['delen'][2]['materiaal_deel']
assert (ROOT/'geluidlab.html').is_file()

# De aanvullende 18-pagina's-GitHub-PDF: acht verschillende practica, niet acht
# extra duplicaatwerkbladen. Paginanummers zijn 1-based; 8 is leeg en 9/12
# bevatten als beeld gerasterde instructietekst (handmatig visueel nagekeken).
extra=pymupdf.open(ROOT/'leerstof 6BCW/labocursus ECB tril golven.pdf')
assert len(extra)==18
source_blocks=[
 (1,2,'Inleidend labo',11,(8,(1,))),
 (3,4,'centripetale kracht',12,(8,(2,))),
 (5,6,'periode van een massa aan een veer',13,(9,(1,))),
 (7,8,'onbekende massa',15,(9,(2,))),
 (9,12,'halve periode',14,(9,(3,))),
 (13,14,'periode van de slinger',16,(10,(1,2,3))),
 (15,16,'Resonantie van een massa',17,(11,(1,2))),
 (17,18,'437',20,(13,(1,2,3))),
]
for first,last,keyword,sheet,(labno,parts) in source_blocks:
 body=' '.join(extra[i].get_text() for i in range(first-1,last))
 assert keyword.lower() in body.lower(),(first,last,keyword)
 assert SHEETS[sheet-1][:2]==(labno,parts),(sheet,labno,parts)
assert len(extra[8].get_images()) and len(extra[11].get_images())
assert 'carbonlint' in ' '.join(LABOS[8]['delen'][2]['stappen']).lower()
assert '4π²/g' in LABOS[9]['delen'][2]['verwerking'] and '4π²/rico' in LABOS[9]['delen'][2]['verwerking']
assert '2π/√g' not in LABOS[9]['delen'][2]['verwerking']
assert 'balans' in SHEET_MATERIAL[(8,(2,))][0] and 'géén balans' in SHEET_MATERIAL[(9,(2,))][0]

# Niet alleen de eerste rij maar ELKE aangetroffen tijdtikkertabel: t′ direct vóór v.
for n,parts,_ in SHEETS:
 for j in parts:
  p=LABOS[n-1]['delen'][j-1]
  for tab in [{'kop':p['kop']}]+p.get('extra_tabellen',[]):
   h=tab['kop']
   if 't′ (s)' in h:
    i=h.index('t′ (s)')
    assert i+1<len(h) and ('v' in h[i+1] or '⟨v⟩' in h[i+1]),(n,j,h)
    assert not any('Δv' in x or 'Δt' in x for x in h),(n,j,h)
assert all(len(LABOS[4]['delen'][k]['extra_tabellen'][0]['rijen'])>=8 for k in (0,1))
assert 'v₀' in ' '.join(LABOS[1]['delen'][0]['grafieken'])
assert len(ANSWERS)>=170
# Verifieer ook de data tegen de ECHTE 5BCW-DOCX, niet alleen tegen de Python-constante.
source=Document(ROOT/'Labocursus_5BCW_LEERLING.docx')
assert any([[c.text for c in row.cells] for row in t.rows]==RUBRIC for t in source.tables), '5BCW-rubric wijkt af van de referentie'

for label in ('LEERLING','LEERKRACHT'):
 docx=ROOT/f'Labocursus_6BCW_{label}.docx';pdf=ROOT/f'Labocursus_6BCW_{label}.pdf'
 d=Document(docx)
 headings=[p.text for p in d.paragraphs if p.style.name.startswith('Heading 1') and re.match(r'^LABO \d{2} –',p.text)]
 assert len(headings)==25,headings
 for i,(_,_,title) in enumerate(SHEETS,1):assert headings[i-1]==f'LABO {i:02d} – {title}'
 assert len(d.inline_shapes)>=24
 text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for t in d.tables for row in t.rows)
 assert 'BRONPRACTICA' not in text.upper() and 'oorspronkelijk practicum' not in text.lower()
 assert 'Bron: Labocursus' not in text and 'Soundry' not in text
 assert 'blz. 79' not in text and 'theorie-PDF' not in text
 assert 'LUCHTBEL' in text.upper() and 'KLASOVERZICHT a(F)' in text and 'KLASOVERZICHT a(m)' in text
 assert 'v₀' in text and '−v₀/a' in text and 't′' in text
 assert 'A. Wat gebeurt er met een wagen' in text and 'Vanaf labo 03' in text
 assert 'CARBONLINT' in text and '4π²/rico' in text and 'Mach-kegel' in text
 assert 'TWINTIG rondjes' in text and 'krijtstreepjes' in text and 'motorkop' in text
 assert '1.1 De OVUR-methode' in text and '1.4 Het verslag' in text and '1.5 Overzicht' in text
 assert 'Blanco_verslag_6BCW_CONCEPT.xlsx' in text and 'De leerkracht bevestigt' in text
 assert sum(p.text=='Context' and p.style.name=='Heading 3' for p in d.paragraphs)==25
 assert sum(p.text.startswith('Toepassen: wat zeggen') for p in d.paragraphs)==25
 assert sum(p.text=='Oriëntatievragen (vóór je begint)' for p in d.paragraphs)==25
 assert len(d.element.xpath('//m:oMathPara'))==len(FORMULAS), 'Formules moeten bewerkbare OfficeMath blijven'
 assert all((n,parts) in SHEET_SAFETY for n,parts in [(1,(1,)),(1,(2,)),(3,(1,)),(3,(2,3)),(4,(1,)),(4,(2,)),(8,(1,)),(8,(2,)),(9,(1,)),(9,(2,)),(9,(3,))])
 assert '6 V-wisselspanningsvoeding' not in ' '.join(SHEET_SAFETY[(1,(1,))])
 for t in d.tables:
  if not t.rows:continue
  header=[c.text for c in t.rows[0].cells]
  if 't′ (s)' in header:
   i=header.index('t′ (s)');assert i+1<len(header) and ('v' in header[i+1] or '⟨v⟩' in header[i+1]),header
 rubric=[[[c.text for c in row.cells] for row in t.rows] for t in d.tables if t.rows and [c.text for c in t.rows[0].cells]==RUBRIC[0]]
 assert len(rubric)==1 and rubric[0]==RUBRIC,'De 5BCW-rubric is aangepast'
 if label=='LEERLING':
  assert 'modelantwoord' not in text.lower() and 'LEERKRACHTENINFO' not in text
  assert 'Mogelijke SMART-vraag; niet vooraf uitdelen' not in text
 else:
  assert text.lower().count('modelantwoord')>70
  assert text.count('Mogelijke SMART-vraag; niet vooraf uitdelen')==23
 with ZipFile(docx) as z:
  assert not any('theorie%20' in z.read(x).decode().lower() for x in z.namelist() if x.endswith('.rels'))
 pdfdoc=pymupdf.open(pdf)
 assert len(pdfdoc)>=95
 assert all(len(p.get_text())>100 for p in pdfdoc)
 assert 'Inhoud' not in pdfdoc[0].get_text() and 'Inhoud' in pdfdoc[1].get_text(), 'Titelpagina is niet apart'
 assert 'g = 4π²/rico' in ' '.join(p.get_text() for p in pdfdoc), 'Correcte slingergrafiekformule moet ook in PDF staan'
 assert all(f in ' '.join(p.get_text() for p in pdfdoc) for f in FORMULA_DISPLAY.values()), 'OMML-tekstequivalent ontbreekt in PDF'
 assert len({img[0] for page in pdfdoc for img in page.get_images()})>=18
 assert any(p.rect.width>p.rect.height for p in pdfdoc)
 pdftext=' '.join(p.get_text() for p in pdfdoc)
 assert 'LABO 25' in pdftext and 'Bijlage F' in pdftext
 assert ('modelantwoord' not in pdftext.lower()) if label=='LEERLING' else ('modelantwoord' in pdftext.lower())
 print(f'{label}: {len(headings)} labo’s, {len(d.tables)} tabellen, {len(FORMULAS)} OfficeMath-formules, {len(d.inline_shapes)} figuren, PDF {len(pdfdoc)} pagina’s; rubric letterlijk.')

report=load_workbook(ROOT/'Blanco_verslag_6BCW_CONCEPT.xlsx')
assert report.active['B2'].value.startswith('TOEGEPASTE FYSICA')
assert report.active['B3'].value.startswith('Niet identiek')
print('OK — acht aanvullende proeven in acht zelfstandige werkbladen gecontroleerd; 5BCW-Exceltemplate en rest van 118-pagina-audit blijven open')
