# -*- coding: utf-8 -*-
"""Opmaaklaag voor de labocursus (docx): koppen, kaders, tabellen, invulvakken,
formules (OMML), figuren en leerkracht-versie-logica.

Alles wordt met echte Word-kopstijlen opgebouwd:
  Heading 1 = Deel / LABO / Bijlage
  Heading 2 = O-V-U-R-fase en paragrafen van Deel 1
  Heading 3 = onderdelen binnen een fase
"""
import os
import re
import html
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE as RT

import omml as _omml

# ---------------------------------------------------------------- constanten
FULL_W = 9638          # bruikbare breedte in twips (A4 met marge 2 cm)
FILL = {
    "doel":      ("E8F5E9", "2E7D32"),   # groen   – lesdoelen
    "lp":        ("FFF8E1", "B25B00"),   # geel    – leerplandoelen
    "voor":      ("F3E5F5", "6A1B9A"),   # paars   – voorkennis
    "veilig":    ("FDECEA", "C62828"),   # rood    – veiligheid
    "stem":      ("E0F2F1", "00695C"),   # groen-blauw – STEM-ontwerp
    "uitleg":    ("E8F1FA", "1F4E79"),   # blauw   – uitleg
    "eval":      ("EEEEEE", "555555"),   # grijs   – evaluatie
    "lkr":       ("FFF0F0", "B71C1C"),   # roze    – leerkrachteninfo
    "antw":      ("FFF3F3", "B71C1C"),   # roze    – modelantwoord
    "deadline":  ("FDECEA", "C62828"),   # rood    – deadline
}
BORDER_KADER = ("9DB6CF", 6)
BORDER_VAK = ("7F7F7F", 6)
BORDER_TABEL = ("999999", 4)
HEAD_FILL = "D9E2F3"
GREY = "555555"

FIGDIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figs")


# ---------------------------------------------------------------- document
def new_document(footer_text="Labocursus Fysica 6BCW – mechanica, trillingen en golven"):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    pf = st.paragraph_format
    pf.space_after = Pt(4)
    pf.line_spacing = 1.08

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(1.0)

    for name, size, color in (("Heading 1", 20, "1F4E79"),
                              ("Heading 2", 15, "1F4E79"),
                              ("Heading 3", 12.5, "2E7D32")):
        s = doc.styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.font.italic = False

    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(footer_text + " – pagina ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GREY)
    _add_page_field(p)
    return doc


def _add_page_field(par):
    r = par.add_run()
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GREY)
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    r._r.append(fld)
    r2 = par.add_run()
    r2.font.size = Pt(8)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " PAGE "
    r2._r.append(it)
    r3 = par.add_run()
    r3.font.size = Pt(8)
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r3._r.append(fld2)


def landscape_section(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21.0)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    return sec


def portrait_section(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    return sec


# ---------------------------------------------------------------- tekst
def _runs(par, text, size=10.5, bold=False, italic=False, color=None):
    """Zet tekst om in runs; **vet** en *cursief* worden gehonoreerd."""
    for part in re.split(r"(\*\*.+?\*\*|(?<!\*)\*(?!\*)[^*]+?\*(?!\*))", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = par.add_run(part[1:-1])
            run.italic = True
        else:
            run = par.add_run(part)
        run.font.size = Pt(size)
        run.font.name = "Calibri"
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return par


def p(doc, text="", size=10.5, bold=False, italic=False, color=None,
      align=None, space_after=None, style=None, indent=None):
    par = doc.add_paragraph(style=style)
    if align == "c":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "r":
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "j":
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    if indent:
        par.paragraph_format.left_indent = Cm(indent)
    _runs(par, text, size=size, bold=bold, italic=italic, color=color)
    return par


def hyperlink(doc, text, url):
    """Klikbare Word-hyperlink naar de theorie-PDF in GitHub."""
    par = doc.add_paragraph()
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    el = OxmlElement('w:hyperlink')
    el.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '1F4E79'); rp.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rp.append(u)
    run.append(rp)
    t = OxmlElement('w:t'); t.text = text; run.append(t)
    el.append(run); par._p.append(el)
    par.paragraph_format.space_after = Pt(3)
    return par


def h1(doc, text):
    par = doc.add_paragraph(text, style="Heading 1")
    par.paragraph_format.space_before = Pt(14)
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.keep_with_next = True
    return par


def h2(doc, text):
    par = doc.add_paragraph(text, style="Heading 2")
    par.paragraph_format.space_before = Pt(10)
    par.paragraph_format.space_after = Pt(3)
    par.paragraph_format.keep_with_next = True
    return par


def h3(doc, text):
    par = doc.add_paragraph(text, style="Heading 3")
    par.paragraph_format.space_before = Pt(8)
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.keep_with_next = True
    return par


def pagebreak(doc):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    run = par.add_run()
    run.add_break(WD_BREAK.PAGE)


def bullets(doc, items, size=10.5, style="List Bullet"):
    for it in items:
        par = doc.add_paragraph(style=style)
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.left_indent = Cm(0.6)
        _runs(par, it, size=size)


def numbered(doc, items, size=10.5):
    for it in items:
        par = doc.add_paragraph(style="List Number")
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.left_indent = Cm(0.6)
        _runs(par, it, size=size)


def vraag(doc, nr, text, size=10.5):
    """Genummerde vraag met vette vraagzin."""
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.left_indent = Cm(0.6)
    r = par.add_run("%s. " % nr)
    r.bold = True
    r.font.size = Pt(size)
    _runs(par, text, size=size)
    return par


# ---------------------------------------------------------------- tabellen
def _set_tbl_borders(tbl, color, sz):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + tag)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _grid(tbl, widths):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = tbl._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w)))
        grid.append(gc)


def tabel(doc, headers, rows, widths=None, size=10, header_fill=HEAD_FILL,
          align=None, first_col_bold=False, min_row_twips=330):
    """Tabel met rasterlijnen, lichtblauwe kopregel en vaste kolombreedtes."""
    ncol = len(headers)
    if widths is None:
        widths = [FULL_W / ncol] * ncol
    tbl = doc.add_table(rows=0, cols=ncol)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    _set_tbl_borders(tbl, *BORDER_TABEL)
    _grid(tbl, widths)

    hr = tbl.add_row()
    # Bij een paginabreuk moet Word de kopregel opnieuw tonen; nooit alleen
    # een voortgezette meettabel zonder kolomnaam en eenheid.
    header_tag = OxmlElement('w:tblHeader')
    header_tag.set(qn('w:val'),'true')
    hr._tr.get_or_add_trPr().append(header_tag)
    for i, htxt in enumerate(headers):
        cell = hr.cells[i]
        cell.width = Twips(int(widths[i]))
        _shade(cell, header_fill)
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.space_before = Pt(1)
        _runs(par, htxt, size=size, bold=True)

    for r in rows:
        row = tbl.add_row()
        # Ruimte voor met de hand in te vullen meetwaarden: ook lege Word-cellen
        # moeten bij afdrukken minstens ongeveer 5,8 mm hoog blijven.
        row.height = Twips(min_row_twips)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i, val in enumerate(r):
            cell = row.cells[i]
            cell.width = Twips(int(widths[i]))
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.space_before = Pt(1)
            _runs(par, "" if val is None else str(val), size=size,
                  bold=(i == 0 and first_col_bold))
    if align:
        for row in tbl.rows:
            for i, a in enumerate(align):
                if a and i < len(row.cells):
                    for par in row.cells[i].paragraphs:
                        par.alignment = {"c": WD_ALIGN_PARAGRAPH.CENTER,
                                         "r": WD_ALIGN_PARAGRAPH.RIGHT,
                                         "l": WD_ALIGN_PARAGRAPH.LEFT}[a]
    return tbl


def tabel_zonder_rand(doc, ncols, widths=None):
    """Tabel zonder randen – bv. om twee figuren naast elkaar te zetten."""
    if widths is None:
        widths = [FULL_W / ncols] * ncols
    tbl = doc.add_table(rows=0, cols=ncols)
    tbl.autofit = False
    _grid(tbl, widths)
    return tbl


# ---------------------------------------------------------------- kaders
def kader(doc, title, items, soort="doel", size=10.5, width=FULL_W,
          with_title=True):
    """Gekleurd kader als 1-cel-tabel met gekleurde titel."""
    fill, color = FILL[soort]
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    _set_tbl_borders(tbl, *BORDER_KADER)
    _grid(tbl, [width])
    row = tbl.rows[0]
    row.height = None
    cell = row.cells[0]
    cell.width = Twips(int(width))
    _shade(cell, fill)
    first = True
    if with_title and title:
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(3)
        _runs(par, title, size=size, bold=True, color=color)
        first = False
    for it in items:
        if isinstance(it, tuple) and it and it[0] == "p":
            par = cell.add_paragraph()
            par.paragraph_format.space_after = Pt(3)
            _runs(par, it[1], size=size,
                  color=color if soort in ('lkr', 'antw') else None)
        elif isinstance(it, tuple) and it and it[0] == "raw":
            par = cell.add_paragraph()
            par.paragraph_format.space_after = Pt(3)
            par._p.append(parse_xml(_omml.omml_inline(it[1])))
        else:
            par = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.left_indent = Cm(0.5)
            r = par.add_run("\u2022 ")
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor.from_string(color)
            _runs(par, it, size=size)
    return tbl


def invul(doc, regels=2, width=FULL_W):
    """Omrand invulvak met vaste hoogte (geen puntjeslijnen)."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    _set_tbl_borders(tbl, *BORDER_VAK)
    _grid(tbl, [width])
    row = tbl.rows[0]
    row.height = Twips(int(180 + 300 * max(1, regels)))
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = row.cells[0]
    cell.width = Twips(int(width))
    par = cell.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    return tbl


def invul_lijn(doc, label=None, width=FULL_W):
    """Eén smal invulvak (bv. voor een naam of een antwoord van één regel)."""
    return invul(doc, regels=1, width=width)


# ---------------------------------------------------------------- formules
def formule(doc, src, size=11):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(6)
    par._p.append(parse_xml(_omml.omml(src)))
    return par


def formule_in(par, src):
    par._p.append(parse_xml(_omml.omml_inline(src)))
    return par


# ---------------------------------------------------------------- figuren
def figuur(doc, bestand, bijschrift=None, breedte_cm=10.0, bron=None):
    pad = os.path.join(FIGDIR, bestand)
    if not os.path.exists(pad):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run("[figuur ontbreekt: %s]" % bestand)
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string("999999")
        return None
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(2)
    par.add_run().add_picture(pad, width=Cm(breedte_cm))
    if bijschrift:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        txt = bijschrift if bron is None else "%s (bron: %s)" % (bijschrift, bron)
        _runs(cap, txt, size=8.5, italic=True, color="777777")
    return par


# ---------------------------------------------------------------- TOC
def toc(doc, fallback, levels="1-2"):
    par = doc.add_paragraph()
    r = par.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    r._r.append(fld)
    r2 = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = ' TOC \\o "%s" \\h \\z \\u ' % levels
    r2._r.append(it)
    r3 = par.add_run()
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "separate")
    r3._r.append(fld2)
    r4 = par.add_run("Rechtermuisknop → 'Veld bijwerken' om de inhoudsopgave te vernieuwen.")
    r4.italic = True
    r4.font.size = Pt(9)
    r4.font.color.rgb = RGBColor.from_string(GREY)
    r5 = par.add_run()
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "end")
    r5._r.append(fld3)
    doc.add_paragraph()
    for line in fallback:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.left_indent = Cm(float(line[0]) * 0.6)
        _runs(par, line[1], size=10)


# ---------------------------------------------------------------- hulp
def leerkracht(T, fn, *a, **kw):
    """Voer fn enkel uit in de leerkrachtenversie."""
    if T:
        return fn(*a, **kw)
    return None


def modelantwoord(doc, tekst, label="Modelantwoord"):
    kader(doc, "%s:" % label, [("p", tekst)], soort="antw", size=10)


def teacher_note(doc, items, title="LEERKRACHTENINFO – didactische tips, verwachte resultaten, valkuilen"):
    kader(doc, title, items, soort="lkr", size=10)
