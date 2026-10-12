# -*- coding: utf-8 -*-
"""Voorlopig invulbaar verslag, totdat het echte 5BCW-Excelbestand beschikbaar is.

Geen claim van letterlijke gelijkheid: de 5BCW-cursus verwijst naar een blanco
Excelbestand op Smartschool; dat bestand zelf staat niet in deze repository.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

OUT=Path(__file__).resolve().parent.parent/'Blanco_verslag_6BCW_CONCEPT.xlsx'
DARK='173859'; BLUE='DDEEF4'; PALE='F2F7FA'
w=Workbook(); s=w.active;s.title='Verslag'
s.sheet_view.showGridLines=False
s.column_dimensions['A'].width=4
for c in 'BCDEFGH':s.column_dimensions[c].width=18
s.column_dimensions['B'].width=24
s.column_dimensions['H'].width=26
s.merge_cells('B2:H2');s['B2']='TOEGEPASTE FYSICA 6BCW · BLANCO VERSLAG (CONCEPT)'
s['B2'].font=Font(size=16,bold=True,color='FFFFFF');s['B2'].fill=PatternFill('solid',fgColor=DARK)
s.row_dimensions[2].height=34
s.merge_cells('B3:H3');s['B3']='Niet identiek aan het originele 5BCW-Excelbestand op Smartschool; vervang zodra dat beschikbaar is.'
s['B3'].font=Font(italic=True,color='9A4B10',size=10)
for row,label in [(5,'Labo en titel'),(6,'Groepsleden / klas'),(7,'Datum en leerkracht')]:
 s.cell(row,2,label).font=Font(bold=True,color=DARK)
 s.merge_cells(start_row=row,start_column=3,end_row=row,end_column=8)
 s.cell(row,3).fill=PatternFill('solid',fgColor=PALE)

def section(row,title):
 s.merge_cells(start_row=row,start_column=2,end_row=row,end_column=8)
 cell=s.cell(row,2,title);cell.fill=PatternFill('solid',fgColor=DARK);cell.font=Font(bold=True,color='FFFFFF',size=11)
 s.row_dimensions[row].height=24

def blank(row,label,lines=2):
 s.cell(row,2,label).font=Font(bold=True,color=DARK)
 s.merge_cells(start_row=row,start_column=3,end_row=row+lines-1,end_column=8)
 cell=s.cell(row,3);cell.fill=PatternFill('solid',fgColor=PALE);cell.alignment=Alignment(vertical='top',wrap_text=True)
 for r in range(row,row+lines):s.row_dimensions[r].height=22

section(9,'1 · VERKENNEN EN ONDERZOEKEN  / 2')
blank(10,'Doel',2);blank(12,'Materiaal',2);blank(14,'Herhaalbare werkwijze',4)
blank(18,'SMART-onderzoeksvraag',3);blank(21,'Hypothese (vooraf)',3);blank(24,'Schema / foto: verwijzing',2)
section(27,'2 · UITVOEREN EN VERWERKEN  / 3')
s.merge_cells('B28:H28');s['B28']='Noteer hieronder ruwe gegevens met eigen grootheidsnamen en eenheden. Kopieer de rijen voor meerdere reeksen.'
s['B28'].alignment=Alignment(wrap_text=True)
for col in range(2,9):
 c=s.cell(29,col);c.fill=PatternFill('solid',fgColor=BLUE);c.font=Font(bold=True,color=DARK)
 c.value=f'grootheid + eenheid {col-1}'
for row in range(30,46):
 for col in range(2,9):
  c=s.cell(row,col);c.border=Border(bottom=Side(style='hair',color='B8CEDA'))
blank(47,'Twee voorbeeldberekeningen per soort',4)
blank(51,'Grafiek(en): punten, assen, eenheden, titel, trendlijn + vergelijking',4)
section(56,'3 · CONCLUDEREN  / 3')
blank(57,'Antwoord op vraag met meetgegevens',4);blank(61,'Theorie en hypothese',3)
blank(64,'Onzekerheden en reflectie',4)
section(69,'4 · ATTITUDE EN VERZORGING (productafspraken)  / 2')
blank(70,'Bestandsvorm en termijn',3)
s['B71']='Aantoonbaar groepsproduct';s['B72']='Feedbackblok behouden';blank(73,'Feedback leerkracht (niet verwijderen)',5)
s.merge_cells('B79:H79');s['B79']='Controle: t′ direct links van v in tijdtikkertabellen; maak een XY-grafiek v(t′).'
s['B79'].font=Font(bold=True,color=DARK)
section(81,'PRODUCTPUNTEN /10 · in te vullen door de leerkracht')
for row,label,maxpunt in [(82,'1. Verkennen & onderzoeken',2),(83,'2. Onderzoek uitvoeren',3),
                            (84,'3. Concluderen',3),(85,'4. Attitude en verzorging (product)',2)]:
 s.cell(row,2,label);s.cell(row,4,maxpunt);s.cell(row,5).fill=PatternFill('solid',fgColor=BLUE)
s['D86']='Totaal /10';s['E86']='=SUM(E82:E85)'
s.merge_cells('B87:H87');s['B87']='Product 60 % (groepspunt); proces/attitude 40 % (individueel volgens Bijlage C).'
s.merge_cells('B88:H88');s['B88']='Indienen via Smartschool: één week na het labo, uiterlijk 23.59 u; concrete datum/zone vooraf door de leerkracht.'
s.freeze_panes='B30'
s.sheet_properties.pageSetUpPr.fitToPage=True
s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A4;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=3
s.print_options.horizontalCentered=True;s.print_area='B2:H88';s.page_margins=PageMargins(left=.3,right=.3,top=.5,bottom=.5,header=.2,footer=.2)
w.save(OUT)
print(OUT)
